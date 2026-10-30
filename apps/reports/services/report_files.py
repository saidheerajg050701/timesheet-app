"""Project hours downloads (Change 11, D-163): the Excel workbook and the CSV of the Detail.

Layout only — every number comes from the ProjectReport (services/project_report/) that the page
shows, so the files always match it. Text cells are spreadsheet-safe (safe_cell, D-68); hours
are numbers with 2 decimals; reasons are never included. Each download writes one EXPORT
audit row with the filters and the row count, never the data.
"""

import csv
import io
from decimal import Decimal

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.audit.models import AuditLog
from apps.audit.services.log import record
from apps.payroll.services.export import safe_cell
from apps.projects.services.tasks import TASK_NOT_RECORDED
from apps.reports.services.project_report import FOOTNOTE, hours
from apps.timesheets.services.calendar import company_tz
from apps.timesheets.services.validation import CATEGORY_LABELS

DETAIL_COLUMNS = (  # (heading, value, numeric) — one row per counted entry, no reasons
    ("date", lambda r: r.work_date.isoformat(), False),
    ("week_start", lambda r: r.week_start.isoformat(), False),
    ("employee_code", lambda r: r.person.employee.employee_code, False),
    ("employee_name", lambda r: r.person.name, False),
    ("team", lambda r: r.person.team, False),
    ("department", lambda r: r.person.department, False),
    ("manager", lambda r: r.person.manager, False),
    ("project_code", lambda r: r.project.code if r.project else "", False),
    ("project_name", lambda r: r.project.name if r.project else "", False),
    ("task_code", lambda r: r.task.code if r.task else "", False),
    ("task_name", lambda r: _task_name(r), False),
    ("charge_code", lambda r: r.charge_code, False),
    ("category", lambda r: CATEGORY_LABELS[r.category], False),
    ("hours", lambda r: hours(r.minutes), True),
    ("week_status", lambda r: r.status, False),
)
MONEY_FORMAT = '"$"#,##0.00'
MANAGING_TABS = (
    "Summary",
    "By project",
    "By team",
    "By manager",
    "By employee",
    "By task",
    "Project × task",
    "Detail",
)
OWN_TABS = ("Summary", "By project", "By task", "Project × task", "By week", "Detail")


def managing_tabs(report):
    """A manager's tabs: By team / By manager only where they change something, as on the
    page (D-165) — never a one-row repeat."""
    skip = {"By team": "team", "By manager": "manager"}
    return [t for t in MANAGING_TABS if t not in skip or skip[t] in report.filters.groups]


def _task_name(row):
    if row.task is not None:
        return row.task.name
    return TASK_NOT_RECORDED if row.project is not None else ""


def file_name(report, extension):
    """project-hours_<from>_to_<to>.<ext>; for an employee's own report
    project-hours_<employee code>_<from>_to_<to>.<ext>."""
    f = report.filters
    who = f"{report.viewer.employee_code}_" if report.own else ""
    return f"project-hours_{who}{f.start.isoformat()}_to_{f.end.isoformat()}.{extension}"


def audit_download(report, fmt, rows):
    """One EXPORT audit row per download: who, which filters, how many rows — never the data."""
    f = report.filters
    return record(
        AuditLog.Action.EXPORT,
        "ProjectReport",
        f"{f.start.isoformat()}..{f.end.isoformat()}",
        actor=report.viewer,
        after={"format": fmt, "filters": f.query(not report.own), "rows": rows},
    )


def detail_csv(report):
    """The Detail as CSV text (the caller adds the byte-order mark, as for payroll CSVs)."""
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow([name for name, _, _ in DETAIL_COLUMNS])
    for r in report.detail:
        writer.writerow(
            [
                f"{value(r):.2f}" if numeric else safe_cell(value(r))
                for _, value, numeric in DETAIL_COLUMNS
            ]
        )
    return out.getvalue()


# --- Excel ------------------------------------------------------------------------------------


def _cell(value):
    return safe_cell(value) if isinstance(value, str) else value


def sheet(book, title, header, rows, hour_columns=()):
    """A tab: bold frozen header, hours as 0.00 numbers, columns sized to their content.
    Written row by row (write-only workbook), so a year of detail stays quick (D-164)."""
    tab = book.create_sheet(title)
    tab.freeze_panes = "A2"
    for index, heading in enumerate(header, start=1):  # widths go before the first row
        values = [heading, *(row[index - 1] for row in rows[:2000])]  # enough to size it
        width = max(len(f"${v:,.2f}" if isinstance(v, float | Decimal) else str(v)) for v in values)
        tab.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 8), 60)
    bold = Font(bold=True)
    head = []
    for heading in header:
        cell = WriteOnlyCell(tab, value=heading)
        cell.font = bold
        head.append(cell)
    tab.append(head)
    hourly = {i for i, heading in enumerate(header) if heading in hour_columns}
    for row in rows:
        cells = []
        for i, value in enumerate(row):
            if i in hourly and isinstance(value, float):  # hours; counts stay whole
                cell = WriteOnlyCell(tab, value=value)
                cell.number_format = "0.00"
                cells.append(cell)
            elif isinstance(value, Decimal):  # money: exact cents, shown as USD (D-171)
                cell = WriteOnlyCell(tab, value=value)
                cell.number_format = MONEY_FORMAT
                cells.append(cell)
            else:
                cells.append(_cell(value))
        tab.append(cells)
    return tab


def _summary(report):
    viewer = report.viewer
    generated = timezone.now().astimezone(company_tz())
    rows = [
        ("Report", "Project hours"),
        ("Generated at", generated.strftime("%Y-%m-%d %H:%M %Z")),
        ("Generated by", f"{viewer.first_name} {viewer.last_name} ({viewer.employee_code})"),
        *report.filters.described(not report.own),
        ("Project hours", hours(report.project_total)),
        ("Non-project time (not filtered by project)", hours(report.non_project_total)),
        ("Total time", hours(report.total)),
        ("People", f"{report.with_hours} with hours of {len(report.people)} people"),
        ("Projects", len(report.by_project)),
        ("Note", "Break and leave are never counted. " + FOOTNOTE),
    ]
    return ["Item", "Value"], rows


def _by_project(report):
    rows = [(p.code, p.name, hours(m), n) for p, m, n in report.by_project]
    rows.append(("Total", "", hours(report.project_total), report.with_hours))
    return ["I-code", "Project", "Hours", "People"], rows


def _groups(groups):
    return [
        (g.label, g.people, hours(g.project), hours(g.non_project), hours(g.total)) for g in groups
    ]


GROUP_HEADER = ["Team", "People", "Project hours", "Non-project time", "Total time"]
HOUR_HEADINGS = {
    "Hours",
    "Project hours",
    "Non-project time",
    "Total time",
    "Own group: project hours",
    "Own group: non-project time",
    "Own group: total time",
    "Including everyone below: project hours",
    "Including everyone below: non-project time",
    "Including everyone below: total time",
    "Meetings",
    "Training",
    "Other",
    "Leave (not counted)",
    "hours",
}


def _by_manager(report):
    header = [
        "Manager",
        "Level",
        "People",
        "Own group: project hours",
        "Own group: non-project time",
        "Own group: total time",
        "Including everyone below: project hours",
        "Including everyone below: non-project time",
        "Including everyone below: total time",
    ]
    rows = [
        (
            g.label,
            g.depth + 1 if g.employee else "",
            g.people,
            hours(g.project),
            hours(g.non_project),
            hours(g.total),
            hours(g.rollup_project),
            hours(g.rollup_non_project),
            hours(g.rollup_total),
        )
        for g in report.by_manager
    ]
    return header, rows


def _task_headings(report):
    return [t.code if t else TASK_NOT_RECORDED for t in report.task_columns]


def _by_employee(report):
    tasks = _task_headings(report)
    header = [
        "Employee code",
        "Name",
        "Status",
        "Team",
        "Manager",
        *tasks,
        "Project hours",
        "Meetings",
        "Training",
        "Other",
        "Non-project time",
        "Total time",
        "Leave (not counted)",
        "Note",
    ]
    rows = []
    for p in report.shown_people:
        rows.append(
            (
                p.employee.employee_code,
                p.name,
                "Inactive" if p.inactive else "Active",
                p.team,
                p.manager,
                *(hours(p.by_task.get(t.code if t else "", 0)) for t in report.task_columns),
                hours(p.project),
                *(hours(p.non_project.get(c, 0)) for c in ("MEETING", "TRAINING", "OTHER")),
                hours(p.non_project_total),
                hours(p.total),
                hours(p.leave),
                p.note,
            )
        )
    return header, rows, {*HOUR_HEADINGS, *tasks}


def _by_task(report):
    rows = [(code, name, hours(m)) for code, name, m in report.by_task]
    rows.append(("Total", "", hours(report.project_total)))
    return ["Task code", "Task", "Hours"], rows


def _matrix(report):
    tasks = _task_headings(report)
    rows = [
        (p.code, p.name, *(hours(m) for m in report.matrix_row(p)), hours(total))
        for p, total, _ in report.by_project
    ]
    rows.append(
        ("Total", "", *(hours(m) for _, _, m in report.by_task), hours(report.project_total))
    )
    return ["I-code", "Project", *tasks, "Total"], rows, {*tasks, "Total"}


def _by_week(report):
    rows = [(w.isoformat(), hours(p), hours(n), hours(p + n)) for w, p, n in report.by_week]
    return ["Week starting", "Project hours", "Non-project time", "Total time"], rows


def _detail(report):
    header = [name for name, _, _ in DETAIL_COLUMNS]
    rows = [tuple(value(r) for _, value, _ in DETAIL_COLUMNS) for r in report.detail]
    return header, rows


def workbook(report):
    """The .xlsx bytes: the tabs that fit the viewer's report (D-163)."""
    book = Workbook(write_only=True)
    tabs = OWN_TABS if report.own else managing_tabs(report)
    for title in tabs:
        if title == "Summary":
            sheet(book, title, *_summary(report), ("Value",))
        elif title == "By project":
            sheet(book, title, *_by_project(report), HOUR_HEADINGS)
        elif title == "By team":
            sheet(book, title, GROUP_HEADER, _groups(report.by_team), HOUR_HEADINGS)
        elif title == "By manager":
            sheet(book, title, *_by_manager(report), HOUR_HEADINGS)
        elif title == "By employee":
            sheet(book, title, *_by_employee(report))
        elif title == "By task":
            sheet(book, title, *_by_task(report), HOUR_HEADINGS)
        elif title == "Project × task":
            sheet(book, title, *_matrix(report))
        elif title == "By week":
            sheet(book, title, *_by_week(report), HOUR_HEADINGS)
        else:
            sheet(book, title, *_detail(report), HOUR_HEADINGS)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
