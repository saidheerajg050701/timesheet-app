import re
from datetime import datetime, time, timedelta
from pathlib import Path

import pytest
from django.urls import reverse
from playwright.sync_api import expect

from audit.models import AuditLog
from directory.models import Employee
from tests.e2e.conftest import E2E, PERIOD, THIS_WEEK, TZ, log_in
from tests.factories import (
    TEST_PASSWORD,
    EmployeeFactory,
    ProjectFactory,
    TimeEntryFactory,
    UserFactory,
)
from timesheets.models import ChangeRequest, PeriodApproval, TimeEntry, Timesheet
from timesheets.services.approvals import approve_period

pytestmark = E2E
NEW_PASSWORD = "Brand-New-Pass-2026"
WEEK2 = PERIOD + timedelta(weeks=1)
MONDAY = THIS_WEEK + timedelta(days=1)  # weeks start on Sunday (D-111)


def week_url(week):
    return reverse("timesheet_week", args=[week])


def status_of(page):
    return page.locator("#week-status")


PORTAL = "I-001 Customer Portal · DEV Development"
BREAK = "Break (unpaid)"


def cell(page, row, day):
    """A grid cell's button, e.g. cell(page, PORTAL, date(2026, 10, 19)) (D-106)."""
    name = re.escape(f"{row}, {day:%a %b %-d}:").replace("/", r"\/")  # Playwright: /…/ regex
    return page.get_by_role("button", name=re.compile(f"^{name}"))


def open_cell(page, row, day):
    cell(page, row, day).click()
    panel = page.locator("dialog[open]")
    expect(panel).to_be_visible()
    return panel


def pick_project(panel, project, typed=None):
    """Search the dialog's project box (its I-code unless told otherwise) and click the
    project, e.g. pick_project(panel, "I-002 Mobile App") (D-157)."""
    panel.get_by_label("Project", exact=True).fill(typed or project.split()[0])
    panel.get_by_role("option", name=project).click()


def add_row(page, category, project=None, task="DEV Development"):
    """+ Add row; project work also chooses the project, then one of its tasks (D-152)."""
    page.get_by_role("button", name="+ Add row").click()
    panel = page.locator("dialog[open]")
    panel.get_by_label("Category").select_option(label=category)
    if project:
        pick_project(panel, project)
        panel.get_by_label("Task").select_option(label=task)
    panel.get_by_role("button", name="Add row", exact=True).click()


# Is the element fully inside the window and the top element at its left, middle and right
# (not covered by another row, not clipped by the scrolling table or the card)?
ON_TOP = """(el) => {
  const r = el.getBoundingClientRect();
  const inside = r.left >= 0 && r.top >= 0 && r.right <= innerWidth && r.bottom <= innerHeight;
  const y = r.top + r.height / 2;
  return inside && [r.left + 4, r.left + r.width / 2, r.right - 4].every((x) => {
    const hit = document.elementFromPoint(x, y);
    return hit === el || el.contains(hit);
  });
}"""
MENU_ITEMS = ("Change row", "Delete row")


def row_menu(page, name):
    """Open a row's ⋮ menu with a real click (D-129): both items visible, on top of everything,
    and the menu starts at the ⋮ button. Returns (toggle, menu)."""
    toggle = page.get_by_role("button", name=f"Row actions for {name}")
    label = page.locator("#entries tr.grid-row", has=toggle).locator(".grid-label-text")
    button, text = toggle.bounding_box(), label.bounding_box()
    assert button["x"] >= text["x"] + text["width"] - 1, name  # right of the name …
    assert button["y"] < text["y"] + text["height"], name  # … on its first line, not below
    toggle.click()
    menu = page.locator(".row-actions .dropdown-menu.show")
    expect(menu).to_have_count(1)
    for item in MENU_ITEMS:
        button = menu.get_by_role("button", name=item)
        expect(button).to_be_visible()
        assert button.evaluate(ON_TOP), f"{name}: {item!r} is covered, clipped or off screen"
    assert abs(menu.bounding_box()["x"] - toggle.bounding_box()["x"]) <= 2, name  # aligned
    return toggle, menu


def menu_item(page, name, item):
    """Open the row's menu and click one of its items, both for real; returns the dialog."""
    _, menu = row_menu(page, name)
    menu.get_by_role("button", name=item).click()
    return page.locator("dialog[open]")


def test_signup_refused_then_accepted(e2e_org, anonymous):
    """T-146: E2E sign-up refused for an unknown email, accepted for a directory employee."""
    page = anonymous
    page.goto("/signup/")
    page.get_by_label("Email").fill("nobody@yourcompany.com")
    page.get_by_label("Password", exact=True).fill(NEW_PASSWORD)
    page.get_by_label("Confirm password").fill(NEW_PASSWORD)
    page.get_by_role("button", name="Sign up").click()
    expect(page.get_by_text("not in the company employee directory")).to_be_visible()

    page.goto("/signup/")
    page.get_by_label("Email").fill("QA1@yourcompany.com")
    page.get_by_label("Password", exact=True).fill(NEW_PASSWORD)
    page.get_by_label("Confirm password").fill(NEW_PASSWORD)
    page.get_by_role("button", name="Sign up").click()
    expect(page.get_by_text("Welcome, Ravi! Your account is ready.")).to_be_visible()
    expect(page.locator("nav.navbar")).to_contain_text("Ravi Kumar")
    expect(page.locator("nav.navbar")).to_contain_text("QA1")


def test_save_draft_then_submit_before_the_deadline(e2e_org, open_as):
    """T-147: E2E save a draft with two slots around a break, reopen it, then submit in time."""
    page = open_as(e2e_org.dev1)
    page.goto(week_url(THIS_WEEK))
    expect(status_of(page)).to_have_text("Not started")
    add_row(page, "Project work", "I-001 Customer Portal")
    panel = open_cell(page, PORTAL, MONDAY)
    panel.get_by_label("Start").fill("08:00")
    panel.get_by_label("End").fill("13:00")
    panel.get_by_label("Reason").fill("Login page")
    panel.get_by_role("button", name="Done").click()
    expect(cell(page, PORTAL, MONDAY)).to_have_text("5.00")
    add_row(page, "Break time")
    panel = open_cell(page, BREAK, MONDAY)
    panel.get_by_label("Start").fill("13:00")
    panel.get_by_label("End").fill("14:00")
    panel.get_by_label("Reason").fill("Lunch")
    panel.get_by_role("button", name="Done").click()

    add_row(page, "Project work", "I-001 Customer Portal")  # again: points at its row (D-127)
    expect(page.locator("[data-week-grid]").get_by_role("status")).to_contain_text(
        "already has a row"
    )
    panel = open_cell(page, PORTAL, MONDAY)  # lists the slot and offers an empty one
    expect(panel.locator(".grid-slot")).to_have_count(2)
    expect(panel.get_by_label("Start").nth(1)).to_be_focused()
    panel.get_by_label("Start").nth(1).fill("14:00")
    panel.get_by_label("End").nth(1).fill("17:00")
    panel.get_by_label("Reason").nth(1).fill("Code review")
    panel.get_by_role("button", name="Done").click()
    expect(cell(page, PORTAL, MONDAY)).to_have_text("8.00 · 2 slots")
    expect(page.locator("#entries [data-role=week-total]")).to_have_text("8.00")  # Break not in it
    page.get_by_role("button", name="Save draft").click()
    expect(status_of(page)).to_have_text("Draft")
    assert TimeEntry.objects.count() == 3

    page.goto(week_url(THIS_WEEK))  # reopened: both slots are still there
    expect(cell(page, PORTAL, MONDAY)).to_have_text("8.00 · 2 slots")
    expect(cell(page, PORTAL, MONDAY)).to_have_attribute("title", "08:00–13:00, 14:00–17:00")
    panel = open_cell(page, PORTAL, MONDAY)
    expect(panel.get_by_label("Reason").first).to_have_value("Login page")
    expect(panel.get_by_label("Reason").nth(1)).to_have_value("Code review")
    panel.get_by_role("button", name="Done").click()
    expect(cell(page, PORTAL, MONDAY)).to_have_text("8.00 · 2 slots")  # the offered slot dropped
    page.get_by_role("button", name="Submit").click()
    expect(status_of(page)).to_have_text("Submitted")
    expect(page.get_by_role("button", name="Save draft")).to_have_count(0)
    expect(page.locator("#entries input")).to_have_count(0)  # read-only now
    expect(cell(page, PORTAL, MONDAY)).to_have_text("8.00 · 2 slots")
    panel = open_cell(page, PORTAL, MONDAY)
    expect(panel).to_contain_text("Login page")
    expect(panel).to_contain_text("14:00")


def test_week_locks_after_the_deadline(e2e_org, make_week, open_as, clock):
    """T-148: E2E an unsubmitted week is locked from Saturday 12:00 (fixed clock)."""
    make_week(e2e_org.dev1, THIS_WEEK, status="DRAFT")
    page = open_as(e2e_org.dev1)
    page.goto(week_url(THIS_WEEK))
    expect(page.get_by_role("button", name="Submit")).to_be_visible()

    clock.move_to(datetime(2026, 10, 24, 12, 0, tzinfo=TZ))  # Saturday 12:00
    page.goto(week_url(THIS_WEEK))
    expect(status_of(page)).to_have_text("Locked")
    expect(page.get_by_role("button", name="Submit")).to_have_count(0)
    expect(page.get_by_role("button", name="Save draft")).to_have_count(0)
    expect(page.locator("#entries input")).to_have_count(0)
    expect(page.locator("#entries")).to_contain_text("Feature work")


def test_change_request_then_manager_edits_directly(e2e_org, make_week, open_as):
    """T-149: E2E change request → manager Edit directly → employee sees it; audit row kept."""
    week = THIS_WEEK - timedelta(weeks=1)
    make_week(e2e_org.dev1, week)
    employee = open_as(e2e_org.dev1)
    employee.goto(week_url(week))
    employee.get_by_text("Request change").click()
    employee.get_by_label("What needs to change, and why?").fill("Monday ran until 18:00")
    employee.get_by_role("button", name="Send to my manager").click()
    expect(employee.locator("#change-request")).to_contain_text("Waiting for your manager")

    manager = open_as(e2e_org.manager)
    manager.goto(reverse("request_inbox"))
    manager.locator("#open-requests").get_by_role("link").first.click()
    manager.get_by_role("link", name="Edit directly").click()
    panel = open_cell(manager, PORTAL, week + timedelta(days=1))  # Monday
    panel.get_by_label("End").first.fill("18:00")  # the saved slot; an empty one is offered too
    panel.get_by_role("button", name="Done").click()
    expect(cell(manager, PORTAL, week + timedelta(days=1))).to_have_text("9.00")
    manager.get_by_label("Note to the employee (optional)").fill("Fixed Monday")
    manager.get_by_role("button", name="Save and submit for the employee").click()
    expect(manager.get_by_role("alert")).to_have_text(
        "Week edited and submitted on the employee's behalf."
    )

    employee.goto(week_url(week))
    expect(status_of(employee)).to_have_text("Submitted")
    expect(open_cell(employee, PORTAL, week + timedelta(days=1))).to_contain_text("18:00")
    change = ChangeRequest.objects.get(timesheet__employee=e2e_org.dev1)
    row = AuditLog.objects.get(action="UPDATE", change_request=change)
    assert row.actor == e2e_org.manager
    assert "18:00" in str(row.after) and "17:00" in str(row.before)


def test_approve_then_payroll_sees_the_records(e2e_org, make_week, open_as):
    """T-150: E2E manager approves a period; the payroll user sees both week records — once
    everyone in the period is approved (Change 9: no hours while the register is OPEN)."""
    make_week(e2e_org.dev1, PERIOD, monday_end=time(19))  # 42 h: 2 h overtime
    make_week(e2e_org.dev1, WEEK2)
    for person in Employee.objects.exclude(pk=e2e_org.dev1.pk):  # the others, 0.00 weeks
        approver = person.manager or e2e_org.hr
        approve_period(approver, person, PERIOD, now=datetime.now(TZ), confirm_unsubmitted=True)
    payroll = open_as(e2e_org.payroll)
    payroll.goto(f"{reverse('payroll')}?period={PERIOD.isoformat()}")
    expect(payroll.locator("#payroll-records")).to_have_count(0)  # OPEN: John still pending

    manager = open_as(e2e_org.manager)
    manager.goto(reverse("approval_detail", args=[PERIOD, "DEV1"]))
    manager.get_by_role("button", name="Approve both weeks").click()
    expect(manager.get_by_role("alert")).to_have_text("Approved John's period.")

    payroll.reload()
    records = payroll.locator("#payroll-records tbody tr", has_text="John Kim")
    expect(records).to_have_count(2)
    expect(records.filter(has_text="Sep 27")).to_contain_text("40.00")
    expect(records.filter(has_text="Sep 27")).to_contain_text("2.00")  # overtime
    expect(records.filter(has_text="Oct 4")).to_contain_text("Current")


def test_reject_with_comment_then_resubmit(e2e_org, make_week, open_as):
    """T-151: E2E reject with a comment; the employee sees it and resubmits both weeks."""
    make_week(e2e_org.dev1, PERIOD)
    make_week(e2e_org.dev1, WEEK2)
    manager = open_as(e2e_org.manager)
    detail = reverse("approval_detail", args=[PERIOD, "DEV1"])
    manager.goto(detail)
    manager.get_by_label("Comment for the employee (required to reject)").fill("Add Friday's demo")
    manager.get_by_role("button", name="Reject").click()
    expect(manager.get_by_role("alert")).to_have_text("Rejected — John can fix and resubmit.")

    employee = open_as(e2e_org.dev1)
    for week in (PERIOD, WEEK2):  # both weeks reopened, with the comment
        employee.goto(week_url(week))
        expect(status_of(employee)).to_have_text("Reopened")
        expect(employee.locator("#approval-status")).to_contain_text("Add Friday's demo")
    for week in (PERIOD, WEEK2):  # a resubmission sets the period back to pending
        employee.goto(week_url(week))
        if week == PERIOD:
            panel = open_cell(employee, PORTAL, PERIOD + timedelta(days=5))  # Friday
            panel.get_by_label("Reason").first.fill("Feature work and demo")
            panel.get_by_role("button", name="Done").click()
        employee.get_by_role("button", name="Submit").click()
        expect(status_of(employee)).to_have_text("Submitted")

    manager.goto(detail)
    expect(manager.get_by_role("button", name="Approve both weeks")).to_be_enabled()


def test_payroll_export_and_mark_paid(e2e_org, make_week, open_as, clock):
    """T-152: E2E payroll finalizes & exports the period's register, downloads the detailed
    CSV, then marks it paid (D-135). Everyone else is approved too: since Change 9 a register
    is ready only then (D-142) — their never-submitted weeks with 0.00 hours."""
    make_week(e2e_org.dev1, PERIOD)
    make_week(e2e_org.dev1, WEEK2)
    for person in Employee.objects.all():  # publishes after each commit
        approver = person.manager or e2e_org.hr
        approve_period(approver, person, PERIOD, now=datetime.now(TZ), confirm_unsubmitted=True)
    page = open_as(e2e_org.payroll)
    page.goto(reverse("payroll_period", args=[PERIOD]))
    expect(page.locator("#register-status")).to_have_text("Ready")
    page.get_by_role("button", name="Finalize & export").click()
    page.locator("#finalize-dialog").get_by_role("button", name="Finalize & export").click()
    expect(page.get_by_role("alert")).to_contain_text("exported as batch EXP-0001")
    with page.expect_download() as download:
        page.get_by_role("link", name="Download detailed CSV").click()
    assert download.value.suggested_filename == (
        "payroll-detail_2026-09-27_to_2026-10-10_EXP-0001.csv"
    )
    text = Path(download.value.path()).read_text(encoding="utf-8-sig")
    assert text.startswith("batch,employee_code") and text.count("DEV1") == 2

    expect(page.locator("#payroll-records")).to_contain_text("Exported")
    page.get_by_role("button", name="Mark as paid").click()
    expect(page.locator("#register-status")).to_have_text("Paid")
    john = page.locator("#payroll-records tbody tr", has_text="John Kim")
    expect(john.filter(has_text="Paid")).to_have_count(2)


def test_hr_deactivates_an_employee(e2e_org, open_as, anonymous):
    """T-153: E2E HR deactivates an employee, who is logged out and cannot log in again."""
    employee = open_as(e2e_org.dev2)
    hr = open_as(e2e_org.hr)
    hr.goto("/admin/directory/employee/")
    hr.locator("tr", has_text="DEV2").get_by_role("checkbox").check()
    hr.locator('select[name="action"]').select_option("deactivate")
    hr.get_by_role("button", name="Go").click()
    expect(hr.get_by_label("End date (last working day)")).to_have_value("2026-10-21")
    hr.get_by_role("button", name="Deactivate").click()
    expect(hr.get_by_text("Deactivated 1 employee(s).")).to_be_visible()

    employee.goto("/")
    expect(employee).to_have_url(re.compile(r"/login/"))
    log_in(anonymous, e2e_org.dev2.email, TEST_PASSWORD)
    expect(anonymous.get_by_text("Wrong email or password.")).to_be_visible()


def test_hr_password_reset_forces_a_change(e2e_org, open_as, anonymous):
    """T-154: E2E HR sets a temporary password; the employee must change it at next login."""
    temporary = "Temporary-Pass-2026"
    hr = open_as(e2e_org.hr)
    hr.goto(f"/admin/directory/employee/{e2e_org.dev1.pk}/change/")
    hr.get_by_role("link", name="Reset password").click()
    hr.get_by_label("Temporary password:", exact=True).fill(temporary)
    hr.get_by_label("Temporary password (again):").fill(temporary)
    hr.get_by_role("button", name="Set temporary password").click()
    expect(hr.get_by_text("They must choose a new one at their next login")).to_be_visible()

    page = anonymous
    log_in(page, e2e_org.dev1.email, temporary)
    expect(page).to_have_url(re.compile(r"/profile/$"))
    expect(page.get_by_text("Choose a new password before you continue.")).to_be_visible()
    page.goto(reverse("timesheets"))
    expect(page).to_have_url(re.compile(r"/profile/$"))
    page.get_by_label("Current password").fill(temporary)
    page.get_by_label("New password", exact=True).fill(NEW_PASSWORD)
    page.get_by_label("Confirm new password").fill(NEW_PASSWORD)
    page.get_by_role("button", name="Change password").click()
    expect(page.get_by_text("Your password was changed.")).to_be_visible()
    page.goto(reverse("timesheets"))
    expect(page).to_have_url(re.compile(r"/timesheets/$"))


def test_forbidden_pages(e2e_org, open_as):
    """T-155: E2E 403 pages for an employee, a payroll-only user and a manager outside a tree."""
    cases = {
        e2e_org.dev1: ["/team/", "/approvals/", "/payroll/", "/admin/"],
        e2e_org.payroll: ["/team/", "/approvals/"],
        e2e_org.manager: [reverse("team_member", args=["OPS1"])],
    }
    for person, urls in cases.items():
        page = open_as(person)
        for url in urls:
            response = page.goto(url)
            assert response.status == 403, (person, url)
            expect(page.get_by_role("heading", name="Not allowed")).to_be_visible()
    assert not PeriodApproval.objects.exists()


def test_weekly_grid_in_the_browser(e2e_org, make_week, open_as):
    """T-167: E2E grid: add rows, two slots by keyboard, live totals, unsaved warning, narrow."""
    page = open_as(e2e_org.dev1)
    page.goto(week_url(THIS_WEEK))
    add_row(page, "Project work", "I-001 Customer Portal")
    add_row(page, "Project work", "I-001 Customer Portal")  # no second row (D-127, T-183)
    expect(page.locator("dialog[open]")).to_have_count(0)
    expect(page.locator('#entries tr[data-row-key^="PROJECT:"]')).to_have_count(1)

    tuesday = THIS_WEEK + timedelta(days=2)
    button = cell(page, PORTAL, tuesday)
    button.focus()
    page.keyboard.press("Enter")  # cells are buttons: the keyboard opens them
    panel = page.locator("dialog[open]")
    panel.get_by_label("Start").fill("08:00")
    panel.get_by_label("End").fill("12:00")
    panel.get_by_label("Reason").fill("Bug fixes")
    expect(button).to_have_text("4.00")  # live while typing
    panel.get_by_role("button", name="+ Add another time slot").click()
    panel.get_by_label("Start").nth(1).fill("13:00")
    panel.get_by_label("End").nth(1).fill("17:30")
    panel.get_by_label("Reason").nth(1).fill("Code review")
    page.keyboard.press("Escape")  # Esc closes; focus goes back to the cell
    expect(page.locator("dialog[open]")).to_have_count(0)
    expect(button).to_be_focused()
    expect(button).to_contain_text("8.50")
    expect(page.locator("#entries [data-role=week-total]")).to_have_text("8.50")

    add_row(page, "PTO")
    panel = open_cell(page, "PTO", tuesday + timedelta(days=1))
    panel.get_by_label("Hours").fill("8")
    panel.get_by_label("Reason").fill("Vacation")
    panel.get_by_role("button", name="Done").click()
    expect(page.locator("#entries [data-role=week-total]")).to_have_text("8.50")  # leave not in it

    warnings = []
    page.on("dialog", lambda dialog: warnings.append(dialog.type))
    page.get_by_role("link", name="Dashboard").click()  # unsaved changes: the browser asks first
    expect(page).to_have_url(re.compile(r"/$"))
    assert "beforeunload" in warnings
    page.goto(week_url(THIS_WEEK))
    assert TimeEntry.objects.count() == 0  # nothing was saved by leaving

    add_row(page, "Project work", "I-001 Customer Portal")
    panel = open_cell(page, PORTAL, tuesday)
    panel.get_by_label("Start").fill("09:00")
    panel.get_by_label("End").fill("12:00")
    panel.get_by_label("Reason").fill("Bug fixes")
    panel.get_by_role("button", name="Done").click()
    page.get_by_role("button", name="Save draft").click()
    expect(status_of(page)).to_have_text("Draft")
    assert TimeEntry.objects.count() == 1

    page.set_viewport_size({"width": 390, "height": 844})
    scroller = page.locator(".grid-scroll")
    assert scroller.evaluate("e => e.scrollWidth > e.clientWidth")  # the grid scrolls sideways
    assert page.evaluate("() => { scrollTo(5000, 0); return scrollX; }") == 0  # the page doesn't
    first = page.locator("#entries tbody th.grid-label").first
    scroller.evaluate("e => e.scrollLeft = 300")
    assert first.evaluate("e => getComputedStyle(e).position") == "sticky"
    box = first.bounding_box()
    assert box["x"] < 20  # still at the left edge while the days scroll


def test_add_row_for_an_existing_row_points_at_it(e2e_org, open_as):
    """T-183: E2E "+ Add row" for a row already there: no error, it is highlighted and focused."""
    page = open_as(e2e_org.dev1)
    page.goto(week_url(THIS_WEEK))
    add_row(page, "Project work", "I-001 Customer Portal")
    add_row(page, "Meetings")
    add_row(page, "Project work", "I-001 Customer Portal")
    expect(page.locator("dialog[open]")).to_have_count(0)  # closed; no error to dismiss
    expect(page.get_by_role("alert")).to_have_count(0)
    expect(page.locator("[data-week-grid]").get_by_role("status")).to_have_text(
        "I-001 Customer Portal · DEV Development already has a row. Click a day to add another "
        "time slot (for example 2:00 PM – 5:00 PM)."
    )
    rows = page.locator("#entries tr.grid-row")
    expect(rows).to_have_count(2)
    expect(page.locator('#entries tr[data-row-key^="PROJECT:"]')).to_have_count(1)
    portal = page.locator('#entries tr[data-row-key^="PROJECT:"]')  # the I-001 row
    expect(portal).to_have_class(re.compile(r"\bis-highlighted\b"))
    expect(cell(page, PORTAL, THIS_WEEK)).to_be_focused()  # its first day, Sunday
    expect(portal).not_to_have_class(re.compile(r"\bis-highlighted\b"), timeout=5000)  # brief

    add_row(page, "PTO")
    add_row(page, "PTO")  # leave rows get the leave wording
    expect(page.locator("[data-week-grid]").get_by_role("status")).to_have_text(
        "PTO already has a row. Click a day to enter or change its hours."
    )
    expect(cell(page, "PTO", THIS_WEEK)).to_be_focused()
    expect(rows).to_have_count(3)
    assert TimeEntry.objects.count() == 0  # nothing saved by any of this


def test_time_without_a_reason(e2e_org, open_as):
    """T-190: E2E time entered without a reason saves and submits; no empty Reason shown."""
    page = open_as(e2e_org.dev1)
    page.goto(week_url(THIS_WEEK))
    add_row(page, "Project work", "I-001 Customer Portal")
    panel = open_cell(page, PORTAL, MONDAY)
    expect(panel.get_by_label("Reason")).to_be_visible()  # still offered, just optional
    panel.get_by_label("Start").fill("09:00")
    panel.get_by_label("End").fill("17:00")
    panel.get_by_role("button", name="Done").click()
    expect(cell(page, PORTAL, MONDAY)).to_have_text("8.00")
    page.get_by_role("button", name="Save draft").click()
    expect(status_of(page)).to_have_text("Draft")
    expect(page.locator("#errors")).to_have_count(0)
    assert TimeEntry.objects.get().reason == ""

    page.get_by_role("button", name="Submit").click()
    expect(status_of(page)).to_have_text("Submitted")
    panel = open_cell(page, PORTAL, MONDAY)
    expect(panel).to_contain_text("09:00")
    expect(panel).to_contain_text("17:00")
    expect(panel.get_by_role("columnheader", name="Reason")).to_have_count(0)


def test_change_and_delete_rows(e2e_org, open_as):
    """T-200: E2E move a wrong row to another project, then delete a row; totals follow."""
    ProjectFactory(number=2, name="Mobile App")
    page = open_as(e2e_org.dev1)
    page.goto(week_url(THIS_WEEK))
    total = page.locator("#entries [data-role=week-total]")
    add_row(page, "Project work", "I-001 Customer Portal")  # the wrong project
    panel = open_cell(page, PORTAL, MONDAY)
    panel.get_by_label("Start").fill("09:00")
    panel.get_by_label("End").fill("12:00")
    panel.get_by_label("Reason").fill("Mobile login")
    panel.get_by_role("button", name="Done").click()
    expect(total).to_have_text("3.00")

    panel = menu_item(page, "I-001 Customer Portal", "Change row")
    expect(panel.get_by_label("Category")).to_have_value("PROJECT")
    expect(panel.get_by_label("Category")).to_be_focused()
    pick_project(panel, "I-002 Mobile App")
    panel.get_by_role("button", name="Change row").click()
    expect(page.locator("dialog[open]")).to_have_count(0)
    mobile = "I-002 Mobile App · DEV Development"
    expect(cell(page, mobile, MONDAY)).to_have_text("3.00")
    expect(page.locator('#entries tr[data-row-key^="PROJECT:"]')).to_have_count(1)
    expect(page.locator("[data-week-grid]").get_by_role("status")).to_have_text(
        f"Moved 1 time slot to {mobile}. Save to keep this change."
    )
    expect(page.get_by_role("button", name="Row actions for I-002 Mobile App")).to_be_focused()
    expect(total).to_have_text("3.00")

    add_row(page, "Meetings")
    panel = open_cell(page, "Meetings", MONDAY)
    panel.get_by_label("Start").fill("13:00")
    panel.get_by_label("End").fill("14:00")
    panel.get_by_role("button", name="Done").click()
    expect(total).to_have_text("4.00")
    toggle = page.get_by_role("button", name="Row actions for Meetings")
    panel = menu_item(page, "Meetings", "Delete row")
    expect(panel).to_contain_text(
        "Delete Meetings? This removes 1 time slot (1.00 h) from this week."
    )
    page.keyboard.press("Escape")  # nothing deleted; focus back on the menu button
    expect(page.locator("dialog[open]")).to_have_count(0)
    expect(toggle).to_be_focused()
    expect(total).to_have_text("4.00")
    menu_item(page, "Meetings", "Delete row").get_by_role("button", name="Delete row").click()
    expect(page.locator('#entries tr[data-row-key="MEETING:"]')).to_have_count(0)
    expect(total).to_have_text("3.00")
    expect(page.locator("#day-summary")).to_contain_text("3:00")  # the server's summary too

    page.get_by_role("button", name="Save draft").click()
    expect(status_of(page)).to_have_text("Draft")
    saved = TimeEntry.objects.get()
    assert (saved.project.code, saved.start_time, saved.end_time, saved.reason) == (
        "I-002",
        time(9),
        time(12),
        "Mobile login",
    )


@pytest.mark.parametrize(
    "size",
    [{"width": 1280, "height": 800}, {"width": 390, "height": 844}],
    ids=["laptop", "narrow"],
)
def test_row_menus_in_a_real_browser(e2e_org, open_as, size):
    """T-202: E2E ⋮ menus open on top and unclipped at the ⋮ button (last row, scrolled page);
    their items open dialogs by mouse and keyboard, also after re-draws; Esc returns to ⋮."""
    ProjectFactory(number=2, name="Mobile App")
    page = open_as(e2e_org.dev1)
    page.set_viewport_size(size)
    page.goto(week_url(THIS_WEEK))
    add_row(page, "Project work", "I-001 Customer Portal")
    add_row(page, "Meetings")
    add_row(page, "Break time")
    add_row(page, "PTO")  # the last row
    for name in ("I-001 Customer Portal", "Meetings", "Break (unpaid)", "PTO"):
        toggle, menu = row_menu(page, name)
        page.keyboard.press("Escape")  # the menu closes; focus stays on ⋮
        expect(menu).to_have_count(0)
        expect(toggle).to_be_focused()

    def dialog_opens(name, item, heading):
        dialog = menu_item(page, name, item)
        expect(dialog).to_have_count(1)
        expect(dialog.get_by_role("heading")).to_contain_text(heading)
        page.keyboard.press("Escape")
        expect(page.locator("dialog[open]")).to_have_count(0)
        expect(page.get_by_role("button", name=f"Row actions for {name}")).to_be_focused()

    dialog_opens("I-001 Customer Portal", "Change row", "Change project, task or category")
    panel = open_cell(page, PORTAL, MONDAY)  # a cell edit
    panel.get_by_label("Start").fill("09:00")
    panel.get_by_label("End").fill("12:00")
    panel.get_by_role("button", name="Done").click()
    panel = menu_item(page, "I-001 Customer Portal", "Change row")
    pick_project(panel, "I-002 Mobile App")
    panel.get_by_role("button", name="Change row").click()  # re-draw 1
    expect(page.get_by_role("button", name="Row actions for I-002 Mobile App")).to_be_focused()

    panel = open_cell(page, "Meetings", MONDAY)  # another row, after the re-draw
    panel.get_by_label("Start").fill("13:00")
    panel.get_by_label("End").fill("14:00")
    panel.get_by_role("button", name="Done").click()
    dialog_opens("Meetings", "Delete row", "Delete row")
    row_menu(page, "Break (unpaid)")[1].get_by_role("button", name="Delete row").click()  # empty
    expect(page.locator('#entries tr[data-row-key="BREAK:"]')).to_have_count(0)  # re-draw 2
    dialog_opens("PTO", "Change row", "Change project, task or category")
    add_row(page, "Training")  # a row added after the re-draws
    dialog_opens("Training", "Change row", "Change project, task or category")

    toggle = page.get_by_role("button", name="Row actions for I-002 Mobile App")
    for key, steps, heading in (("Enter", 1, "Change project"), (" ", 2, "Delete row")):
        toggle.focus()
        page.keyboard.press(key)  # opens the menu
        expect(page.locator(".row-actions .dropdown-menu.show")).to_have_count(1)
        for _ in range(steps):
            page.keyboard.press("ArrowDown")
        page.keyboard.press(key)  # the focused item
        expect(page.locator("dialog[open] h2")).to_contain_text(heading)
        page.keyboard.press("Escape")
        expect(toggle).to_be_focused()

    last = page.locator("#entries tr.grid-row").last
    last.evaluate("r => scrollBy(0, r.getBoundingClientRect().bottom - innerHeight + 4)")
    toggle, _ = row_menu(page, "Training")  # the last row at the bottom of the window
    page.keyboard.press("Escape")
    page.locator(".grid-scroll").evaluate("e => e.scrollLeft = e.scrollWidth")  # days scrolled
    row_menu(page, "Meetings")
    page.keyboard.press("Escape")
    expect(page.locator(".dropdown-menu.show")).to_have_count(0)


def test_register_from_last_approval_to_paid(e2e_org, make_week, open_as, clock):
    """T-216: E2E the register stays OPEN until the last approval, is READY right after it,
    then Finalize & export (one batch, downloads) and Mark as paid (D-132–D-135)."""
    clock.move_to(datetime(2026, 10, 12, 9, 0, tzinfo=TZ))  # Monday, before the 12:00 deadline
    people = list(Employee.objects.all())
    for person in people:
        make_week(person, PERIOD)
        make_week(person, WEEK2)
        if person != e2e_org.dev2:
            approver = person.manager or e2e_org.hr  # the CEO: approved by HR (D-11)
            approve_period(approver, person, PERIOD, now=datetime.now(TZ))
    payroll = open_as(e2e_org.payroll)
    payroll.goto(reverse("payroll_period", args=[PERIOD]))
    expect(payroll.locator("#register-status")).to_have_text("Open")
    pending = payroll.locator("#pending-by-manager")
    expect(pending).to_contain_text("Sara Ng")
    expect(pending).to_contain_text("Priya Shah")  # her approver
    expect(payroll.get_by_role("button", name="Finalize & export")).to_have_count(0)  # OPEN

    manager = open_as(e2e_org.manager)
    manager.goto(reverse("approval_detail", args=[PERIOD, "DEV2"]))
    manager.get_by_role("button", name="Approve both weeks").click()
    expect(manager.get_by_role("alert")).to_have_text("Approved Sara's period.")

    payroll.reload()
    expect(payroll.locator("#register-status")).to_have_text("Ready")
    expect(payroll.locator("#ready-reason")).to_contain_text("Everyone who must be approved")
    payroll.get_by_role("button", name="Finalize & export").click()
    payroll.locator("#finalize-dialog").get_by_role("button", name="Finalize & export").click()
    expect(payroll.get_by_role("alert")).to_contain_text("exported as batch EXP-0001")
    expect(payroll.locator("#register-status")).to_have_text("Finalized & exported")
    expect(payroll.locator("#register tbody tr")).to_have_count(len(people))
    with payroll.expect_download() as download:
        payroll.get_by_role("link", name="Download register CSV").click()
    assert download.value.suggested_filename == (
        "payroll-register_2026-09-27_to_2026-10-10_EXP-0001.csv"
    )
    lines = Path(download.value.path()).read_text(encoding="utf-8-sig").splitlines()
    assert lines[0].startswith("employee_code,") and len(lines) == len(people) + 1

    payroll.get_by_role("button", name="Mark as paid").click()
    expect(payroll.locator("#register-status")).to_have_text("Paid")
    expect(payroll.get_by_role("button", name="Mark as paid")).to_have_count(0)
    payroll.goto(reverse("payroll"))
    row = payroll.locator('#pay-periods tr[data-period="2026-09-27"]')
    expect(row).to_contain_text("Paid")
    expect(row).to_contain_text(f"{len(people)} of {len(people)}")
    expect(row).to_contain_text("EXP-0001")


def test_manager_edits_and_approves_a_locked_employee(e2e_org, open_as):
    """T-224: E2E Sara never submitted (both weeks locked): her manager enters her hours in the
    grid, approves through the confirmation — the last approval — and payroll's register turns
    READY and is finalized & exported (D-139–D-142)."""
    for person in Employee.objects.exclude(pk=e2e_org.dev2.pk):  # everyone else: approved
        approver = person.manager or e2e_org.hr
        approve_period(approver, person, PERIOD, now=datetime.now(TZ), confirm_unsubmitted=True)
    payroll = open_as(e2e_org.payroll)
    payroll.goto(reverse("payroll_period", args=[PERIOD]))
    expect(payroll.locator("#register-status")).to_have_text("Open")  # deadline passed: Overdue
    expect(payroll.locator("#pending-by-manager")).to_contain_text("Locked – not submitted")

    manager = open_as(e2e_org.manager)
    manager.goto(reverse("approval_detail", args=[PERIOD, "DEV2"]))
    expect(manager.locator("#week-1")).to_contain_text("Not submitted by employee")
    manager.get_by_role("link", name="Edit week 1").click()
    add_row(manager, "Project work", "I-001 Customer Portal")
    monday = PERIOD + timedelta(days=1)
    panel = open_cell(manager, PORTAL, monday)
    panel.get_by_label("Start").fill("09:00")
    panel.get_by_label("End").fill("17:00")
    panel.get_by_role("button", name="Done").click()
    manager.get_by_role("button", name="Save the week").click()
    expect(manager.get_by_role("alert")).to_contain_text("Week of Sep 27 saved for Sara.")

    manager.get_by_role("button", name="Approve both weeks").click()
    dialog = manager.locator("#approve-dialog")
    expect(dialog).to_contain_text(
        "Week 1 was not submitted by Sara Ng; it will be approved with 8.00 hours."
    )
    expect(dialog).to_contain_text("Week 2 has 0.00 hours.")
    dialog.get_by_role("button", name="Approve both weeks").click()
    expect(manager.get_by_role("alert")).to_have_text("Approved Sara's period.")
    sheet = Timesheet.objects.get(employee=e2e_org.dev2, week_start=PERIOD)
    assert (sheet.status, sheet.submitted_at) == ("APPROVED", None)

    payroll.reload()
    expect(payroll.locator("#register-status")).to_have_text("Ready")
    payroll.get_by_role("button", name="Finalize & export").click()
    payroll.locator("#finalize-dialog").get_by_role("button", name="Finalize & export").click()
    expect(payroll.get_by_role("alert")).to_contain_text("exported as batch EXP-0001")
    sara = payroll.locator('#register tr[data-employee="DEV2"]')
    expect(sara.locator('[data-role="pay-total"]')).to_have_text("8.00")


def test_tasks_and_charge_codes_in_a_real_browser(e2e_org, open_as):
    """T-236: E2E add a Project work row with a task (only the project's tasks offered), enter
    time, change the task with ⋮, save and submit; the manager sees the charge code (D-152)."""
    mobile = ProjectFactory(number=2, name="Mobile App")
    mobile.tasks.set(mobile.tasks.filter(code__in=["DEV", "TST"]))
    page = open_as(e2e_org.dev1)
    page.goto(week_url(THIS_WEEK))
    page.get_by_role("button", name="+ Add row").click()
    panel = page.locator("dialog[open]")
    panel.get_by_label("Category").select_option(label="Project work")
    expect(panel.get_by_label("Task")).to_be_hidden()  # asked once a project is chosen
    pick_project(panel, "I-002 Mobile App")
    expect(panel.get_by_label("Task").locator("option")).to_have_text(
        ["Choose…", "DEV Development", "TST Testing & QA"]
    )
    panel.get_by_role("button", name="Add row", exact=True).click()
    expect(panel.get_by_role("alert")).to_have_text("Choose a task.")
    panel.get_by_label("Task").select_option(label="DEV Development")
    panel.get_by_role("button", name="Add row", exact=True).click()
    dev = "I-002 Mobile App · DEV Development"
    row = page.locator("#entries tr.grid-row", has_text="I-002 Mobile App")
    expect(row.get_by_title("Charge code")).to_have_text("I-002-DEV")
    panel = open_cell(page, dev, MONDAY)
    panel.get_by_label("Start").fill("09:00")
    panel.get_by_label("End").fill("12:00")
    panel.get_by_label("Reason").fill("Test plan")
    panel.get_by_role("button", name="Done").click()

    panel = menu_item(page, "I-002 Mobile App", "Change row")
    expect(panel.get_by_label("Task")).to_have_value("DEV")
    panel.get_by_label("Task").select_option(label="TST Testing & QA")
    panel.get_by_role("button", name="Change row").click()
    tst = "I-002 Mobile App · TST Testing & QA"
    expect(cell(page, tst, MONDAY)).to_have_text("3.00")
    expect(page.locator("#entries").get_by_title("Charge code")).to_have_text(["I-002-TST"])
    page.get_by_role("button", name="Save draft").click()
    expect(status_of(page)).to_have_text("Draft")
    page.get_by_role("button", name="Submit").click()
    expect(status_of(page)).to_have_text("Submitted")
    entry = TimeEntry.objects.get()
    assert (entry.project.code, entry.task.code, entry.reason) == ("I-002", "TST", "Test plan")

    manager = open_as(e2e_org.manager)
    manager.goto(reverse("approval_detail", args=[THIS_WEEK - timedelta(weeks=1), "DEV1"]))
    expect(manager.get_by_title("Charge code")).to_have_text(["I-002-TST"])
    expect(manager.locator("[data-week-grid]").last).to_contain_text(tst)  # week 2 of the period


@pytest.mark.parametrize(
    "size", [{"width": 1280, "height": 900}, {"width": 390, "height": 844}], ids=["laptop", "phone"]
)
def test_project_search_in_a_real_browser(e2e_org, open_as, size):
    """T-243: E2E the project search box: type "portal", pick I-001 with ↓ + Enter, choose a
    task, add the row, enter time, save; ⋮ Change row: "No matching project", then "i-002";
    Esc closes the list before the dialog; on a laptop and a phone (D-157)."""
    ProjectFactory(number=2, name="Mobile App")
    page = open_as(e2e_org.dev1)
    page.set_viewport_size(size)
    page.goto(week_url(THIS_WEEK))
    page.get_by_role("button", name="+ Add row").click()
    panel = page.locator("dialog[open]")
    panel.get_by_label("Category").select_option(label="Project work")
    box = panel.get_by_role("combobox", name="Project")
    expect(box).to_have_attribute("aria-expanded", "false")
    box.click()  # the empty box lists every active project, by I-code
    expect(panel.get_by_role("listbox").get_by_role("option")).to_have_text(
        ["I-001 Customer Portal", "I-002 Mobile App"]
    )
    expect(box).to_have_attribute("aria-expanded", "true")
    box.fill("portal")
    expect(panel.get_by_role("listbox").get_by_role("option")).to_have_text(
        ["I-001 Customer Portal"]
    )
    expect(panel.locator("[data-role=project-count]")).to_have_text("1 project")  # announced
    box.press("ArrowDown")
    expect(box).to_have_attribute("aria-activedescendant", re.compile(r"add-row-project-list-\d+"))
    box.press("Enter")
    expect(box).to_have_value("I-001 Customer Portal")
    expect(panel.get_by_role("listbox")).to_be_hidden()
    panel.get_by_label("Task").select_option(label="DEV Development")
    panel.get_by_role("button", name="Add row", exact=True).click()
    expect(page.locator("dialog[open]")).to_have_count(0)
    panel = open_cell(page, PORTAL, MONDAY)
    panel.get_by_label("Start").fill("09:00")
    panel.get_by_label("End").fill("12:00")
    panel.get_by_role("button", name="Done").click()
    page.get_by_role("button", name="Save draft").click()
    expect(status_of(page)).to_have_text("Draft")

    page.get_by_role("button", name=f"Row actions for {PORTAL}").click()
    page.locator(".row-actions .dropdown-menu.show").get_by_role(
        "button", name="Change row"
    ).click()
    panel = page.locator("dialog[open]")
    box = panel.get_by_role("combobox", name="Project")
    expect(box).to_have_value("I-001 Customer Portal")  # the row's own project
    box.fill("zzz")
    expect(panel.get_by_role("listbox")).to_contain_text("No matching project")
    expect(panel.locator("[data-role=project-count]")).to_have_text("No matching project")
    box.press("Escape")  # closes the list only
    expect(panel.get_by_role("listbox")).to_be_hidden()
    expect(page.locator("dialog[open]")).to_have_count(1)
    box.fill("i-002")
    option = panel.get_by_role("option", name="I-002 Mobile App")
    expect(panel.get_by_role("listbox").get_by_role("option")).to_have_count(1)
    assert option.bounding_box()["width"] <= size["width"]  # fits the phone
    option.click()
    expect(box).to_have_value("I-002 Mobile App")
    expect(panel.get_by_label("Task")).to_have_value("DEV")
    panel.get_by_role("button", name="Change row").click()
    mobile = "I-002 Mobile App · DEV Development"
    expect(cell(page, mobile, MONDAY)).to_have_text("3.00")
    page.get_by_role("button", name="Save draft").click()
    expect(status_of(page)).to_have_text("Draft")
    assert TimeEntry.objects.get().project.code == "I-002"
    assert page.evaluate("document.documentElement.scrollWidth") <= size["width"]


REPORT_FILES = "2026-09-27_to_2026-10-10"  # the last pay period at the e2e clock


def downloads(page, prefix):
    """Download the report's Excel and CSV; returns their file names."""
    names = []
    for link in ("Download Excel", "Download CSV"):
        with page.expect_download() as download:
            page.get_by_role("link", name=link).click()
        names.append(download.value.suggested_filename)
    assert names == [f"{prefix}{REPORT_FILES}.xlsx", f"{prefix}{REPORT_FILES}.csv"]
    return names


def choose_report_project(page, typed, project):
    box = page.get_by_role("combobox", name="Project")
    box.fill(typed)
    page.get_by_role("listbox").get_by_role("option", name=project).click()
    expect(page.locator("[data-role=project-chips]")).to_contain_text(project)
    page.get_by_role("button", name="Apply").click()


def test_employee_project_hours_report(e2e_org, make_week, open_as):
    """T-252: E2E Sara opens Reports → Project hours: only her own hours; filters to one
    project with the search box; downloads Excel and CSV with her code in the names (D-160)."""
    mobile = ProjectFactory(number=2, name="Mobile App")
    week = make_week(e2e_org.dev2, PERIOD, status="APPROVED")
    week.entries.filter(work_date=PERIOD + timedelta(days=5)).update(project=mobile)  # Friday
    make_week(e2e_org.dev1, PERIOD, status="APPROVED")  # John's hours never show for Sara
    page = open_as(e2e_org.dev2)
    page.get_by_role("link", name="Reports").click()
    expect(page.get_by_role("heading", name="Project hours", level=1)).to_be_visible()
    expect(page.locator("#summary-project")).to_have_text("40.00")
    expect(page.locator("#summary-people")).to_have_text("1 with hours of 1 person")
    expect(page.locator("#by-employee")).to_have_count(0)
    expect(page.get_by_label("People", exact=True)).to_have_count(0)  # no scope for an employee
    choose_report_project(page, "mobile", "I-002 Mobile App")
    expect(page).to_have_url(re.compile(r"project=I-002"))
    expect(page.locator("#summary-project")).to_have_text("8.00")
    expect(page.locator("#by-project-table tbody tr")).to_have_count(1)
    downloads(page, "project-hours_DEV2_")
    exports = AuditLog.objects.filter(action="EXPORT", actor=e2e_org.dev2)
    assert [e.after["filters"]["project"] for e in exports] == [["I-002"], ["I-002"]]


def test_director_project_hours_report(e2e_org, make_week, open_as):
    """T-253: E2E a director filters to one project, groups by manager, drills into a
    manager's group and downloads Excel and CSV (D-160, D-162)."""
    director = EmployeeFactory(
        employee_code="DIR", first_name="Dina", last_name="Ross", manager=e2e_org.ceo,
        email="dir@yourcompany.com",
    )  # fmt: skip
    UserFactory(employee=director)
    for manager in (e2e_org.manager, e2e_org.ops_manager):
        manager.manager = director
        manager.save()
    mobile = ProjectFactory(number=2, name="Mobile App")
    make_week(e2e_org.dev1, PERIOD, status="APPROVED")
    make_week(e2e_org.dev2, PERIOD, status="APPROVED")
    ops = make_week(e2e_org.ops1, PERIOD, status="APPROVED")
    ops.entries.update(project=mobile)
    TimeEntryFactory(
        timesheet=ops, work_date=PERIOD + timedelta(days=6), category="MEETING", project=None,
        task=None, start_time=time(9), end_time=time(10), minutes=60,
    )  # fmt: skip
    page = open_as(director)
    page.goto(reverse("project_hours"))
    expect(page.locator("#summary-project")).to_have_text("120.00")
    choose_report_project(page, "portal", "I-001 Customer Portal")
    expect(page.locator("#summary-project")).to_have_text("80.00")
    expect(page.locator("#summary-non-project")).to_have_text("1.00")  # not filtered by project
    page.get_by_label("Group by").select_option(label="Manager")
    page.get_by_role("button", name="Apply").click()
    rows = page.locator("#by-manager-table tbody tr")
    expect(rows).to_have_count(3)
    expect(rows.nth(0)).to_contain_text("Dina Ross")
    expect(page.locator("#by-manager-table tfoot")).to_contain_text("80.00")
    page.locator("#by-manager-table").get_by_role("link", name="Priya Shah").click()
    expect(page).to_have_url(re.compile(r"scope=manager.*manager=MANAGER"))
    expect(page.get_by_label("People", exact=True)).to_have_value("manager")
    expect(page.locator("#summary-project")).to_have_text("80.00")
    people = page.locator("#by-employee-table tbody tr")
    expect(people.first).to_contain_text("John Kim")
    expect(page.locator("#by-employee-table")).not_to_contain_text("Lena Park")
    expect(page.locator("#by-employee-table tr.is-zero")).to_contain_text(["No time entered"] * 2)
    downloads(page, "project-hours_")
    assert AuditLog.objects.filter(action="EXPORT", actor=director).count() == 2


def test_pay_and_project_budget_in_a_real_browser(e2e_org, make_week, open_as):
    """T-270: E2E HR adds a pay change; a director opens Reports → Project budget, sees an Over
    budget project, drills into a manager's group, downloads the costs Excel and the hours CSV;
    HR sees By employee and the Rate missing list (D-166–D-171)."""
    import io
    from decimal import Decimal

    from openpyxl import load_workbook

    from directory.models import PayRecord

    director = EmployeeFactory(
        employee_code="DIR", first_name="Dina", last_name="Ross", manager=e2e_org.ceo,
        email="dir@yourcompany.com",
    )  # fmt: skip
    UserFactory(employee=director)
    for manager in (e2e_org.manager, e2e_org.ops_manager):
        manager.manager = director
        manager.save()
    mobile = ProjectFactory(number=2, name="Mobile App", budget=Decimal("100000"))
    portal = e2e_org.project
    portal.budget = Decimal("2000")
    portal.save()
    make_week(e2e_org.dev1, PERIOD, status="APPROVED")
    make_week(e2e_org.dev2, PERIOD, status="APPROVED")
    make_week(e2e_org.ops1, PERIOD, status="APPROVED").entries.update(project=mobile)
    PayRecord.objects.create(
        employee=e2e_org.dev2, pay_type="HOURLY", hourly_rate=Decimal("40"),
        effective_from=PERIOD, system_actor="test",
    )  # fmt: skip

    hr = open_as(e2e_org.hr)
    hr.goto(reverse("admin:directory_employee_pay", args=[e2e_org.dev1.pk]))
    hr.get_by_label("Salaried").check()
    hr.get_by_label("Annual salary (USD)").fill("104,000")
    expect(hr.locator("#pay-suggestion")).to_contain_text("Suggested: $50.00 an hour")
    hr.locator("#pay-suggestion").get_by_role("button", name="Use it").click()
    expect(hr.get_by_label("Hourly rate (USD)")).to_have_value("50.00")
    hr.get_by_label("Effective from").fill(PERIOD.isoformat())
    hr.get_by_role("button", name="Save pay change").click()
    expect(hr.locator("#pay-current")).to_contain_text("$50.00 an hour")
    assert AuditLog.objects.filter(entity_type="Pay", actor=e2e_org.hr).count() == 1

    page = open_as(director)
    page.get_by_role("button", name="Reports").click()
    page.get_by_role("link", name="Project budget").click()
    portal_row = page.locator('#by-project-table tr[data-project="I-001"]')
    expect(portal_row).to_contain_text("Over budget")  # 80 h: $2,000 + $1,600 > $2,000
    expect(portal_row).to_contain_text("$3,600.00")
    page.get_by_label("Group by").select_option(label="Manager")
    page.get_by_role("button", name="Apply").click()
    page.locator("#by-manager-table").get_by_role("link", name="Priya Shah").click()
    expect(page).to_have_url(re.compile(r"scope=manager.*manager=MANAGER"))
    expect(page.locator("#by-employee-table")).to_contain_text("John Kim")
    expect(page.locator("#by-employee-table")).not_to_contain_text("Lena Park")
    with page.expect_download() as costs:
        page.get_by_role("link", name="Costs Excel").click()
    assert costs.value.suggested_filename == f"project-costs_{REPORT_FILES}.xlsx"
    with page.expect_download() as hours:
        page.get_by_role("link", name="Hours CSV").click()
    assert hours.value.suggested_filename == f"project-hours_{REPORT_FILES}.csv"

    hr.goto(reverse("project_budget"))
    expect(hr.locator("#by-employee-table")).to_contain_text("Lena Park")
    expect(hr.locator("#rate-missing")).to_contain_text("Lena Park")
    expect(hr.locator("#rate-missing-warning")).to_contain_text("40.00 h have no pay rate")
    with hr.expect_download() as download:
        hr.get_by_role("link", name="Costs Excel").click()
    book = load_workbook(io.BytesIO(Path(download.value.path()).read_bytes()))
    assert "By employee" in book.sheetnames and "Rate missing" in book.sheetnames
