"""Browser tests: Reports — Project hours and Project budget, and employee pay."""

import re
from datetime import time, timedelta
from pathlib import Path

from django.urls import reverse
from playwright.sync_api import expect

from apps.audit.models import AuditLog
from tests.e2e.conftest import E2E, PERIOD
from tests.e2e.helpers import (
    REPORT_FILES,
    choose_report_project,
    downloads,
)
from tests.factories import (
    EmployeeFactory,
    ProjectFactory,
    TimeEntryFactory,
    UserFactory,
)

pytestmark = E2E


def test_employee_project_hours_report(e2e_org, make_week, open_as):
    """T-252: E2E Sara opens Reports → Project hours: only her own hours; filters to one
    project with the search box; downloads Excel and CSV with her code in the names (D-160)."""
    mobile = ProjectFactory(number=2, name="Mobile App")
    week = make_week(e2e_org.dev2, PERIOD, status="APPROVED")
    week.entries.filter(work_date=PERIOD + timedelta(days=5)).update(project=mobile)  # Friday
    make_week(e2e_org.dev1, PERIOD, status="APPROVED")  # John's hours never show for Sara
    page = open_as(e2e_org.dev2)
    page.get_by_role("link", name="Reports").click()
    expect(page.get_by_role("heading", name="Project hours", level=1)).to_be_visible()
    expect(page.locator("#summary-project")).to_have_text("40.00")
    expect(page.locator("#summary-people")).to_have_text("1 with hours of 1 person")
    expect(page.locator("#by-employee")).to_have_count(0)
    expect(page.get_by_label("People", exact=True)).to_have_count(0)  # no scope for an employee
    choose_report_project(page, "mobile", "I-002 Mobile App")
    expect(page).to_have_url(re.compile(r"project=I-002"))
    expect(page.locator("#summary-project")).to_have_text("8.00")
    expect(page.locator("#by-project-table tbody tr")).to_have_count(1)
    downloads(page, "project-hours_DEV2_")
    exports = AuditLog.objects.filter(action="EXPORT", actor=e2e_org.dev2)
    assert [e.after["filters"]["project"] for e in exports] == [["I-002"], ["I-002"]]


def test_director_project_hours_report(e2e_org, make_week, open_as):
    """T-253: E2E a director filters to one project, groups by manager, drills into a
    manager's group and downloads Excel and CSV (D-160, D-162)."""
    director = EmployeeFactory(
        employee_code="DIR", first_name="Dina", last_name="Ross", manager=e2e_org.ceo,
        email="dir@yourcompany.com",
    )  # fmt: skip
    UserFactory(employee=director)
    for manager in (e2e_org.manager, e2e_org.ops_manager):
        manager.manager = director
        manager.save()
    mobile = ProjectFactory(number=2, name="Mobile App")
    make_week(e2e_org.dev1, PERIOD, status="APPROVED")
    make_week(e2e_org.dev2, PERIOD, status="APPROVED")
    ops = make_week(e2e_org.ops1, PERIOD, status="APPROVED")
    ops.entries.update(project=mobile)
    TimeEntryFactory(
        timesheet=ops, work_date=PERIOD + timedelta(days=6), category="MEETING", project=None,
        task=None, start_time=time(9), end_time=time(10), minutes=60,
    )  # fmt: skip
    page = open_as(director)
    page.goto(reverse("project_hours"))
    expect(page.locator("#summary-project")).to_have_text("120.00")
    choose_report_project(page, "portal", "I-001 Customer Portal")
    expect(page.locator("#summary-project")).to_have_text("80.00")
    expect(page.locator("#summary-non-project")).to_have_text("1.00")  # not filtered by project
    page.get_by_label("Group by").select_option(label="Manager")
    page.get_by_role("button", name="Apply").click()
    rows = page.locator("#by-manager-table tbody tr")
    expect(rows).to_have_count(3)
    expect(rows.nth(0)).to_contain_text("Dina Ross")
    expect(page.locator("#by-manager-table tfoot")).to_contain_text("80.00")
    page.locator("#by-manager-table").get_by_role("link", name="Priya Shah").click()
    expect(page).to_have_url(re.compile(r"scope=manager.*manager=MANAGER"))
    expect(page.get_by_label("People", exact=True)).to_have_value("manager")
    expect(page.locator("#summary-project")).to_have_text("80.00")
    people = page.locator("#by-employee-table tbody tr")
    expect(people.first).to_contain_text("John Kim")
    expect(page.locator("#by-employee-table")).not_to_contain_text("Lena Park")
    expect(page.locator("#by-employee-table tr.is-zero")).to_contain_text(["No time entered"] * 2)
    downloads(page, "project-hours_")
    assert AuditLog.objects.filter(action="EXPORT", actor=director).count() == 2


def test_pay_and_project_budget_in_a_real_browser(e2e_org, make_week, open_as):
    """T-270: E2E HR adds a pay change; a director opens Reports → Project budget, sees an Over
    budget project, drills into a manager's group, downloads the costs Excel and the hours CSV;
    HR sees By employee and the Rate missing list (D-166–D-171)."""
    import io
    from decimal import Decimal

    from openpyxl import load_workbook

    from apps.compensation.models import PayRecord

    director = EmployeeFactory(
        employee_code="DIR", first_name="Dina", last_name="Ross", manager=e2e_org.ceo,
        email="dir@yourcompany.com",
    )  # fmt: skip
    UserFactory(employee=director)
    for manager in (e2e_org.manager, e2e_org.ops_manager):
        manager.manager = director
        manager.save()
    mobile = ProjectFactory(number=2, name="Mobile App", budget=Decimal("100000"))
    portal = e2e_org.project
    portal.budget = Decimal("2000")
    portal.save()
    make_week(e2e_org.dev1, PERIOD, status="APPROVED")
    make_week(e2e_org.dev2, PERIOD, status="APPROVED")
    make_week(e2e_org.ops1, PERIOD, status="APPROVED").entries.update(project=mobile)
    PayRecord.objects.create(
        employee=e2e_org.dev2, pay_type="HOURLY", hourly_rate=Decimal("40"),
        effective_from=PERIOD, system_actor="test",
    )  # fmt: skip

    hr = open_as(e2e_org.hr)
    hr.goto(reverse("admin:directory_employee_pay", args=[e2e_org.dev1.pk]))
    hr.get_by_label("Salaried").check()
    hr.get_by_label("Annual salary (USD)").fill("104,000")
    expect(hr.locator("#pay-suggestion")).to_contain_text("Suggested: $50.00 an hour")
    hr.locator("#pay-suggestion").get_by_role("button", name="Use it").click()
    expect(hr.get_by_label("Hourly rate (USD)")).to_have_value("50.00")
    hr.get_by_label("Effective from").fill(PERIOD.isoformat())
    hr.get_by_role("button", name="Save pay change").click()
    expect(hr.locator("#pay-current")).to_contain_text("$50.00 an hour")
    assert AuditLog.objects.filter(entity_type="Pay", actor=e2e_org.hr).count() == 1

    page = open_as(director)
    page.get_by_role("button", name="Reports").click()
    page.get_by_role("link", name="Project budget").click()
    portal_row = page.locator('#by-project-table tr[data-project="I-001"]')
    expect(portal_row).to_contain_text("Over budget")  # 80 h: $2,000 + $1,600 > $2,000
    expect(portal_row).to_contain_text("$3,600.00")
    page.get_by_label("Group by").select_option(label="Manager")
    page.get_by_role("button", name="Apply").click()
    page.locator("#by-manager-table").get_by_role("link", name="Priya Shah").click()
    expect(page).to_have_url(re.compile(r"scope=manager.*manager=MANAGER"))
    expect(page.locator("#by-employee-table")).to_contain_text("John Kim")
    expect(page.locator("#by-employee-table")).not_to_contain_text("Lena Park")
    with page.expect_download() as costs:
        page.get_by_role("link", name="Costs Excel").click()
    assert costs.value.suggested_filename == f"project-costs_{REPORT_FILES}.xlsx"
    with page.expect_download() as hours:
        page.get_by_role("link", name="Hours CSV").click()
    assert hours.value.suggested_filename == f"project-hours_{REPORT_FILES}.csv"

    hr.goto(reverse("project_budget"))
    expect(hr.locator("#by-employee-table")).to_contain_text("Lena Park")
    expect(hr.locator("#rate-missing")).to_contain_text("Lena Park")
    expect(hr.locator("#rate-missing-warning")).to_contain_text("40.00 h have no pay rate")
    with hr.expect_download() as download:
        hr.get_by_role("link", name="Costs Excel").click()
    book = load_workbook(io.BytesIO(Path(download.value.path()).read_bytes()))
    assert "By employee" in book.sheetnames and "Rate missing" in book.sheetnames
